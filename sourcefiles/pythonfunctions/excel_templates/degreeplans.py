
import sys
import os
import json

import csv
import random

from openpyxl import load_workbook
from openpyxl import Workbook
# use this to define the function define_start_column
from openpyxl.utils import column_index_from_string

# this assigns cells colors
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter


import time
import os

from openpyxl.styles import Font
from openpyxl.styles import Border, Side, Alignment

import sqlite3



def make_degreeplan_excel_files(
            universityname,
            degreename,schoolname,
            savepath,semesterdictionary,
            schoolnamecolor,bigbordercolor,smallerbordercolor,
            gridlinecolor,rowtextcolor,titlecolor,
            mainbackgroundcolor,paddingbackgroundcolor,
            subheadingbordercolor,
            datafontname="Helvetica",titlefontname='Calibri',
            logofontname="Barlow",
            logocolor="ffffff",
            subheadingsize=22,
            subheadingcolor='ffffff',
            headingsfontname='Calibri',
            headingsfontcolor='ffffff',
            datarowheight=30,
            headingrowheight=40,
            columnscaler=1
):




    '''
    This uses the result of the scrape semester data function to make an excel file out of it.
    For the CSV files I did them fully vertical. For these excel files, since I can control what columns and cells things are going,
    I want it to be more of a horizontal feel. 

    '''

    totalhours=0
    numberofsemesters=len(semesterdictionary)
    excelobject=[]                        
    print(f'Excel object has been reset to {len(excelobject)}\n\n\n')



    '''
    Use Font and OpenPyXL to create nicely formatted tabular data
    '''

    semesterworkbook=Workbook()
    # gets the default worksheet
    ws=semesterworkbook.active



        
    titlefont=Font(size=26,bold=True,color=titlecolor,name=titlefontname)

    headingborder=Border(bottom=Side(style='mediumDashDot',color=smallerbordercolor))

    headingsfonts=Font(size=18,bold=True,color='ffffff')
    subheadingsfont=Font(size=17,bold=True,color='ffffff')
    # lighter burnt orange
    utnamefont=Font(size=19, color=schoolnamecolor, name='Georgia', bold=True)
    # update later
    semesterfont=Font(bold=True,size=18,color='ffffff')

    # applied to the actual semester data. 
    datafont=Font(size=15,name=datafontname,color=rowtextcolor)


    leftalign=Alignment(horizontal='left')
    centeralign=Alignment(horizontal='center')
    



    # total height of 40


    # first row after title that is. Start at row 3 since rows 1 and 2 were merged
    firstrow=[f'{degreename}',"","","",f'{schoolname}',universityname]
        
    # colindex is like i            
    # start= 1 means start at the first column
    for col_index, value in enumerate(firstrow, start=1):
        # adjust row here
        headingcell=ws.cell(row=3, column=col_index, value=value)
        # use.font to assign the font I see
        if col_index==6:

            headingcell.font=utnamefont
            headingcell.border=headingborder
        elif col_index==1:
            headingcell.font=headingsfonts
            headingcell.border=headingborder
            headingcell.alignment=centeralign
        else:
            headingcell.font=headingsfonts
            headingcell.border=headingborder

            # print(datacell.font.size)


    blankrow=['','','','','','','','']
    # appends to the next empty row. Like writer.writerow(['']) for csvs
    ws.append(blankrow)
    
                
    subheadingrow=['','Course Code','Course Name','Hours','Category','Upper/Lower Division']
    ws.append(subheadingrow)
    # now apply styles
    for cell in ws[5]:
        cell.font=subheadingsfont
        cell.alignment=leftalign


    # now the meat of the file, the data
    



    # for each semester
    rowcount=0
    for semesternum in range(numberofsemesters):
        semester=list(semesterdictionary)[semesternum]
        # semester courses is a dictionary of its own as well
        semestercourses=semesterdictionary[semester]
        # that means the semester will be on the first row
        excelobject.append([f'{semester}   '])
        rowcount+=1
        for coursenameindex in range(len(semestercourses)):

            coursename=list(semestercourses)[coursenameindex]
            # if its NOT a list of lists, ie, a normal course entry
            if len(semestercourses[coursename])==4 and not isinstance(semestercourses[coursename][0],list):
                coursecode, coursehours, upperdivstatus, coursecategory=semestercourses[coursename]
                excelobject.append(["",coursecode,coursename,coursehours,coursecategory,upperdivstatus])
                rowcount+=1
                if coursehours!='':
                    totalhours+=int(coursehours)
            
            else:
                listofcourses=semestercourses[coursename]
                for i in range(len(listofcourses)):
                    coursecode, coursehours, upperdivstatus, coursecategory=listofcourses[i]
                    excelobject.append(["",coursecode,coursename,coursehours,coursecategory,upperdivstatus])
                    rowcount+=1

                    if coursehours!='':
                        totalhours+=int(coursehours)


        # line between semesters
        excelobject.append(['','','','',''])
        rowcount+=1

    # adding to excel file
    # rowval is important for adding rows in order
    rowval=6

    for rowentry in range(len(excelobject)):
        # update it here so it updates by row not column...although
        rowval += 1
                # start at column one, and then remember the padding rows are added later. 
        for col_index, value in enumerate(excelobject[rowentry], start=1):
            # change alignnment here. 1 indexed not 0
            datacell = ws.cell(row=rowval, column=col_index, value=value)

            if col_index==1:
                # use.font to assign the font I see
                datacell.font=semesterfont
                datacell.alignment=Alignment(horizontal='right')

            elif col_index in [2,3]:
                datacell.font=datafont
                datacell.alignment=leftalign
            elif col_index==4: #hourscol
                datacell.font=datafont
                datacell.alignment=centeralign
            elif col_index in [5,6]:
                datacell.font=datafont
                datacell.alignment=leftalign
            
    
                


    print(f'len excel object={len(excelobject)}, ie number of datarows')
    

# -----------------------------end data stuff --------------------------------------------------------
    ws.append(blankrow)
    totalhoursrow=['','','',f'Total Hours: {totalhours}','','']

    # this lastrow value is actually used for the last TWO rows
    lastrowindex=rowcount+6
    for col_index, value in enumerate(totalhoursrow, start=1):
        
        # use.font to assign the font I see
        cell=ws.cell(row=lastrowindex+1, column=col_index, value=value)
        # FF=full opacity 
        cell.font=datafont

    
    lastrow=['DegreeView','','','','','degreeviewsite.com']
    lastrowindex+=2 # 6 rows before we start data stuff. Then rowcount is the amount of data. 

    # change the logo colors here
    for col_index, value in enumerate(lastrow, start=1):
        
        # use.font to assign the font I see
        lastcell=ws.cell(row=lastrowindex, column=col_index, value=value)
        # FF=full opacity 
        if col_index==6:
            # site link cell
            lastcell.font=Font(name='Roboto',size=19, bold=True, color='e7e9eb')
            lastcell.alignment=Alignment(horizontal='left',vertical='bottom')

        else:
            # logo cell
            lastcell.font=Font(name=logofontname,size=30, bold=True, color=logocolor)
            lastcell.alignment=Alignment(horizontal='left',vertical='center')

        # one more cause now we wrote the actual last row there
        ws.row_dimensions[lastrowindex+1].height = 50



    # ---------------formatting stuff------------------
    
    # add padding
    
    
    ws.insert_rows(1)      #new row
    ws.insert_cols(1)       # new col

    # ws.columns gets columns. Each column is a tuple of cells
    for column_cells in ws.columns:
        # max function gets the max in a list. Same for min
        # generator must be in ()
        cellwidthgenerator=(len(str(cell.value)) if cell.value else 0 for cell in column_cells)
        colwidth = max(cellwidthgenerator)

        firstcell=column_cells[0]
        # every cell has a .column_letter attribute
        col_letter = firstcell.column_letter
        col_index=column_index_from_string(col_letter)
        # scale the width factor to make the columns wider
        if col_index==7:  
            # make the UT Austin column alot wider
            ws.column_dimensions[col_letter].width = int(colwidth)*2
        elif col_index==6:
            ws.column_dimensions[col_letter].width = int(colwidth)*1.5
        elif col_index==5:
            ws.column_dimensions[col_letter].width=int(colwidth)*.7
        elif col_index==3:
            ws.column_dimensions[col_letter].width=int(colwidth)*2
            
        # the semester column
        elif col_index==2:
            ws.column_dimensions[col_letter].width=int(colwidth)*1.26
    

        else:
            ws.column_dimensions[col_letter].width = int(colwidth)*1.2
        



    # set sizes for the padding row and column. Units of default font which is usually Arial it seems

    ws.column_dimensions['A'].width=15
    ws.column_dimensions['H'].width=15
    ws.row_dimensions[1].height=25
    ws.row_dimensions[lastrowindex+2].height=25

    # ------------------------------------TITLE STUFF-------------------------------------------

    # create a merged 'Degree view and degreename header can't lie".
    ws.merge_cells('B2:G3')    
    ws.row_dimensions[2].height = 30
    ws.row_dimensions[3].height = 30
    mergedrowcontent=f'{degreename} Sample Semester Layout'
    
    # refer to top left of merged cells
    titlecell=ws['B2']
    titlecell.value=mergedrowcontent
    titlecell.alignment = Alignment(horizontal='center', vertical='center')
    titlecell.font=titlefont

# ------------------------ CONTINUE FORMATTING STUFF -----------------------------------------------------
    # apply a border around the entire file -----------------------------------------------------------------------------

    # make it responsive based on the amount of data
    # remove the +1 here to have no border
    rows = list(ws['B2':f'G{lastrowindex+1}'])
    min_row = 2
    max_row = lastrowindex+1

    min_col = column_index_from_string('B')  
    max_col = column_index_from_string('G') 

    
    # where border code used to be
    
    # make the entire worksheet a color:
    paddingfill=PatternFill(fill_type="solid", start_color=paddingbackgroundcolor) #end_color='0000FF' fill_type="gray125" or linear later

    # have the background be like a padding. 
    for row in ws.iter_rows(min_row=1, max_row=lastrowindex+2, min_col=1, max_col=8):
        for cell in row:
            cell.fill = paddingfill
    
    # reverse the background for cells with content:

    backgroundcolor=PatternFill(fill_type="solid", start_color=mainbackgroundcolor)
    # re-add gridline borders:

    gridline_border = Border(
        # make gridline border not exactly white
        left=Side(border_style="thin", color=gridlinecolor),
        right=Side(border_style="thin", color=gridlinecolor),
        top=Side(border_style="thin", color=gridlinecolor),
        bottom=Side(border_style="thin", color=gridlinecolor)
    )

    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.fill = backgroundcolor
            cell.border=gridline_border
    
    # this is the border that will go around everything.
    # Use logic to only add border to the cells on the outside.

    entire_ws_border=Side(style='thick',color=bigbordercolor)
    
    for row_index, row in enumerate(rows, start=min_row):
        for col_index, cell in enumerate(row, start=min_col):
            current = cell.border

            
            # remember we defined side. and you do Border(Side=style))
            '''
            This is the final peice of the puzzle that adds borders.
            If sides if empty, then
            cell.border = Border(), which defines no borders. Normally it will be entire_ws_border border preset. 
            '''
            cell.border = Border(
    top=entire_ws_border if row_index == min_row else current.top,
    bottom=entire_ws_border if row_index == max_row else current.bottom,
    left=entire_ws_border if col_index == min_col else current.left,
    right=entire_ws_border if col_index == max_col else current.right
)

    
    # where is this one?
    # this is like the subheading border
    for row in ws.iter_rows(min_row=6, max_row=6, min_col=3, max_col=7):
        for cell in row:
            current = cell.border
            cell.border = Border(
                top=current.top,
                bottom=Side(style='medium',color=subheadingbordercolor),  # Only change bottom keep thick border around everything
                left=current.left,
                right=current.right
            )

    
    semesterworkbook.save(savepath)
    print(f'Saved workbook at {savepath}')
