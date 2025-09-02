from openpyxl import Workbook
from openpyxl.styles import Border, Side

wb = Workbook()
ws = wb.active

ws['E5'].value = "Test"
thick_side = Side(style='thick', color='000000')
ws['E5'].border = Border(left=thick_side, right=thick_side, top=thick_side, bottom=thick_side)

wb.save("test.xlsx")
